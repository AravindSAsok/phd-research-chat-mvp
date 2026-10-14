import os
from pathlib import Path

import streamlit as st
from dotenv import load_dotenv
from docx import Document
from google import genai
from google.genai import types
from chat_logic import (
    CLARIFICATION_MESSAGE, NO_EVIDENCE_MESSAGE, api_error_message, build_context,
    build_instruction, format_source, generate_with_retries, needs_clarification,
    recent_conversation, replace_source_ids,
)
from pypdf import PdfReader

load_dotenv()

APP_NAME = "p!nk \u2014 The Research Chat"

st.set_page_config(page_title=APP_NAME, layout="wide")
st.title(APP_NAME)
st.caption("A focused space for rigorous research and clear evidence.")

st.markdown(
    """
    <style>
    :root {
        --pink-primary: #F4A6C1;
        --pink-rose: #FCE4EC;
        --page-background: #FFFDFD;
        --chat-background: #FFFFFF;
        --text-primary: #000000;
        --text-secondary: #444444;
        --pink-border: #E8D5DC;
    }

    .stApp,
    [data-testid="stAppViewContainer"] {
        color: var(--text-primary);
        background: var(--page-background);
    }

    [data-testid="stHeader"] {
        background: transparent;
    }

    .main .block-container {
        max-width: 1120px;
        padding-top: 2rem;
        padding-bottom: 3rem;
    }

    [data-testid="stAppViewContainer"] :is(h1, h2, h3, h4, h5, h6, p, span, label, li, small, strong, em, blockquote, td, th, summary),
    [data-testid="stSidebar"] :is(h1, h2, h3, h4, h5, h6, p, span, label, li, small, strong, em, blockquote, td, th, summary),
    [data-testid="stAlert"] :is(p, span, div, a) {
        color: var(--text-primary) !important;
    }

    .stApp h1 {
        color: var(--text-primary) !important;
        font-weight: 650;
        letter-spacing: -0.03em;
    }

    .stApp h1::after {
        content: "  \\2727";
        color: var(--pink-primary);
        font-size: 0.5em;
        vertical-align: 0.2em;
    }

    [data-testid="stCaptionContainer"] p {
        color: var(--text-secondary) !important;
    }

    [data-testid="stCaptionContainer"] p::before {
        content: "\\2727  ";
        color: var(--pink-primary);
    }

    [data-testid="stSidebar"] {
        background: linear-gradient(180deg, #FCE4EC 0%, #FFF8FA 55%, #FFFDFD 100%);
        border-right: 1px solid var(--pink-border);
    }

    [data-testid="stSidebar"] h2,
    [data-testid="stSidebar"] h3 {
        color: var(--text-primary) !important;
    }

    [data-testid="stFileUploader"],
    [data-testid="stFileUploader"] * {
        color: var(--text-primary) !important;
    }

    [data-testid="stFileUploader"] section {
        background: var(--chat-background);
        border: 1px solid var(--pink-border);
        border-radius: 14px;
    }

    .stButton > button,
    [data-testid="stBaseButton-secondary"] {
        color: var(--text-primary) !important;
        background: var(--chat-background);
        border: 1px solid var(--pink-border);
        border-radius: 12px;
        transition: background-color 140ms ease, border-color 140ms ease;
    }

    .stButton > button:hover,
    [data-testid="stBaseButton-secondary"]:hover {
        color: var(--text-primary) !important;
        background: var(--pink-rose);
        border-color: var(--pink-primary);
    }

    .stButton > button[kind="primary"],
    .stButton > button[data-testid="stBaseButton-primary"],
    [data-testid="stBaseButton-primary"] {
        color: var(--text-primary) !important;
        background: var(--pink-primary);
        border-color: var(--pink-primary);
    }

    .stButton > button[kind="primary"]:hover,
    .stButton > button[data-testid="stBaseButton-primary"]:hover,
    [data-testid="stBaseButton-primary"]:hover {
        color: var(--text-primary) !important;
        background: #EE93B2;
        border-color: #EE93B2;
    }

    [data-testid="stChatMessage"] {
        color: var(--text-primary) !important;
        background: var(--chat-background);
        border: 1px solid var(--pink-border);
        border-radius: 16px;
        padding: 0.85rem 1rem;
    }

    [data-testid="stChatMessage"] *,
    [data-testid="stChatMessage"] [data-testid="stMarkdownContainer"] *,
    [data-testid="stChatMessage"] [data-testid="stText"] * {
        color: var(--text-primary) !important;
    }

    [data-testid="stChatMessage"] a {
        color: #222222 !important;
        text-decoration: underline;
        text-decoration-color: var(--pink-primary);
    }

    [data-testid="stChatMessage"] code,
    [data-testid="stChatMessage"] pre {
        color: var(--text-primary) !important;
        background: #FFF8FA;
        border-color: var(--pink-border);
    }

    [data-testid="stChatMessage"] th {
        background: var(--pink-rose);
    }

    [data-testid="stChatMessage"] table,
    [data-testid="stChatMessage"] th,
    [data-testid="stChatMessage"] td {
        border-color: var(--pink-border);
    }

    [data-testid="stChatInput"] {
        color: var(--text-primary) !important;
        background: var(--chat-background) !important;
        border: 1px solid var(--pink-border) !important;
        border-radius: 14px;
        box-shadow: 0 2px 10px rgba(120, 70, 90, 0.06);
    }

    [data-testid="stChatInput"] textarea,
    [data-testid="stChatInputTextArea"] {
        color: #000000 !important;
        -webkit-text-fill-color: #000000 !important;
        caret-color: #000000 !important;
        background: #FFFFFF !important;
    }

    [data-testid="stChatInput"] textarea::placeholder,
    [data-testid="stChatInputTextArea"]::placeholder {
        color: #666666 !important;
        -webkit-text-fill-color: #666666 !important;
        opacity: 1;
    }

    [data-testid="stChatInput"] textarea::selection,
    [data-testid="stChatInputTextArea"]::selection {
        color: #000000;
        background: #FCE4EC;
    }

    button:focus-visible,
    textarea:focus-visible,
    input:focus-visible {
        outline: 3px solid #D97E9F !important;
        outline-offset: 2px;
    }

    hr {
        border-color: var(--pink-border);
    }
    </style>
    """,
    unsafe_allow_html=True,
)

def get_secret(name, default=""):
    try:
        return st.secrets.get(name, default)
    except Exception:
        # st.secrets can raise when no local secrets file exists.
        return default


api_key = os.getenv("GEMINI_API_KEY") or get_secret("GEMINI_API_KEY", "")
model = os.getenv("GEMINI_MODEL") or get_secret("GEMINI_MODEL", "gemini-3.8-flash")


def chunk_text(text, size=7000, overlap=700):
    text = text.strip()
    if not text:
        return []
    chunks = []
    start = 0
    while start < len(text):
        end = min(start + size, len(text))
        chunks.append(text[start:end])
        if end == len(text):
            break
        start = end - overlap
    return chunks


def extract_document(uploaded_file):
    """Return text passages with PDF page metadata when the file has pages."""
    suffix = Path(uploaded_file.name).suffix.lower()
    passages = []

    if suffix == ".pdf":
        reader = PdfReader(uploaded_file)
        for page_number, page in enumerate(reader.pages, start=1):
            page_text = (page.extract_text() or "").strip()
            for chunk in chunk_text(page_text):
                passages.append({"text": chunk, "page": page_number})
        return passages

    if suffix == ".docx":
        doc = Document(uploaded_file)
        parts = [paragraph.text for paragraph in doc.paragraphs if paragraph.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells if cell.text.strip()))
        text = "\n".join(part for part in parts if part.strip())
    elif suffix == ".txt":
        raw = uploaded_file.getvalue()
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = raw.decode("latin-1")
    else:
        return []

    return [{"text": chunk, "page": None} for chunk in chunk_text(text)]


if "documents" not in st.session_state:
    st.session_state.documents = {}
if "messages" not in st.session_state:
    st.session_state.messages = []

with st.sidebar:
    st.markdown(f"**{APP_NAME}**")
    st.header("Your documents")
    uploads = st.file_uploader(
        "Upload PDF, DOCX, or TXT files",
        type=["pdf", "docx", "txt"],
        accept_multiple_files=True,
        help="Files and extracted text are held in this app session; this app does not save them to a database.",
    )
    if st.button("Process uploaded files", type="primary", use_container_width=True):
        if not uploads:
            st.info("Choose at least one file first.")
        else:
            new_docs = {}
            total_chars = 0
            for uploaded in uploads:
                try:
                    passages = extract_document(uploaded)
                except Exception:
                    st.error(f"Could not read {uploaded.name}. Check that it is a valid, unencrypted file.")
                    continue
                if not passages:
                    st.warning(f"No readable text found in {uploaded.name}. Scanned PDFs may need OCR.")
                    continue
                new_docs[uploaded.name] = passages
                total_chars += sum(len(passage["text"]) for passage in passages)
            st.session_state.documents = new_docs
            st.session_state.messages = []
            st.success(f"Processed {len(new_docs)} document(s), {total_chars:,} characters.")

    if st.session_state.documents:
        st.markdown("**Ready to chat:**")
        for name, passages in st.session_state.documents.items():
            st.write(f"{name} - {len(passages)} text passage(s)")
        if st.button("Clear documents and chat", use_container_width=True):
            st.session_state.documents = {}
            st.session_state.messages = []
            st.rerun()
    st.divider()
    st.caption("Document text is sent to Gemini to answer your questions. Review Google's data-use terms before uploading confidential research.")

if not api_key:
    st.warning("Add GEMINI_API_KEY to a local .env file or Streamlit Secrets to start chatting.")
    st.code("GEMINI_API_KEY=your_api_key_here\nGEMINI_MODEL=gemini-3.8-flash", language="text")
    st.stop()

if not st.session_state.documents:
    st.info("Upload one or more research papers in the sidebar, then click **Process uploaded files**.")
    st.markdown("**Try asking:** What are the main findings? Compare the methodologies. What research gaps are mentioned?")
    st.stop()

for message in st.session_state.messages:
    with st.chat_message(message["role"]):
        st.markdown(message["content"])

prompt = st.chat_input("Ask a question about your uploaded research...")
if prompt:
    # Capture completed previous turns before adding the current question.
    history = recent_conversation(st.session_state.messages)
    st.session_state.messages.append({"role": "user", "content": prompt})
    with st.chat_message("user"):
        st.markdown(prompt)

    context, sources = build_context(
        prompt,
        st.session_state.documents,
        history=history,
    )
    failed = False
    with st.chat_message("assistant"):
        if needs_clarification(prompt, history):
            answer = CLARIFICATION_MESSAGE
            st.markdown(answer)
        elif not context:
            answer = NO_EVIDENCE_MESSAGE
            st.markdown(answer)
        else:
            instruction = build_instruction(prompt, context, history)
            with st.spinner("Searching your documents and drafting an answer..."):
                try:
                    # Disable the SDK's default five attempts. The bounded helper below
                    # owns retries so requests are not multiplied and Retry-After is honored.
                    http_options = types.HttpOptions(
                        retry_options=types.HttpRetryOptions(attempts=1)
                    )
                    with genai.Client(api_key=api_key, http_options=http_options) as client:
                        response = generate_with_retries(
                            lambda: client.models.generate_content(
                                model=model,
                                contents=instruction,
                                config=types.GenerateContentConfig(temperature=0.2),
                            )
                        )
                    answer = response.text or "The model returned an empty answer. Please try again."
                    answer = replace_source_ids(answer, sources)
                except Exception as exc:
                    answer = api_error_message(exc)
                    failed = True
            st.markdown(answer)

            if not failed:
                source_list = list(dict.fromkeys(format_source(source) for source in sources.values()))
                if source_list:
                    st.text("Sources from retrieved excerpts:\n" + "\n".join(f"- {source}" for source in source_list))

    # Keep errors visible, but exclude them from future factual conversation context.
    st.session_state.messages.append({
        "role": "assistant",
        "content": answer,
        "error": failed,
    })
